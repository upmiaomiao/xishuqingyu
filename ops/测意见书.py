# -*- coding: utf-8 -*-
"""审核意见书（docx）自检（服务器上跑）。

自检要看的是"打开之后是不是一份能签字的意见书"，所以查的是**内容**：
封面字段、分级标题、每条的审核项/依据/理由/证据页码、人工复核栏、落款栏、档位说明。
只查"文件能打开"是不够的 —— 生成成功但正文空白的 docx 一样能打开。
"""
import os
import sys

from docx import Document

sys.path.insert(0, "/home/test/xishu_qingyu_serve")
AUDIT_HOME = "/data/eia_audit"
sys.path.insert(0, AUDIT_HOME)
from xishu_pipeline.audit_anchor import build_anchors, load_parsed, load_result  # noqa: E402
from xishu_pipeline.audit_docx import SECTIONS, export_docx  # noqa: E402

import glob  # noqa: E402

OUT = os.path.join(AUDIT_HOME, "_审核结果", "导出")
REPORTS = "/data/eia_reports"
OK = FAIL = 0


def check(name, cond, extra=""):
    global OK, FAIL
    if cond:
        OK += 1
        print("  √ %s" % name)
    else:
        FAIL += 1
        print("  × %s %s" % (name, extra))


def pdf_of(name):
    for p in glob.glob(os.path.join(REPORTS, "**", "*.pdf"), recursive=True):
        if os.path.basename(p) == name:
            return p
    return None


target = sys.argv[1] if len(sys.argv) > 1 else None
names = [j[:-5] + ".pdf" for j in sorted(os.listdir(os.path.join(AUDIT_HOME, "_审核结果")))
         if j.endswith(".json")]
for nm in names:
    try:
        res = load_result(nm)
    except FileNotFoundError:
        continue
    if not res.get("file"):
        continue
    fname = res["file"]["name"]
    if target and target not in fname:
        continue
    path = pdf_of(fname)
    if not path:
        continue
    print("=" * 78)
    print(fname[:70])
    parsed = load_parsed(fname, path, allow_parse=False)
    anchors = build_anchors(res, parsed)
    r = export_docx(fname, res, anchors, OUT, {})
    check("生成 ok", r.get("ok"), str(r)[:150])
    if not r.get("ok"):
        continue
    doc = Document(r["path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    all_txt = text + "\n" + cells

    check("封面标题", "环境影响报告书审核意见书" in text)
    check("含项目名称行", "项目名称" in cells)
    check("含报告文件行", "报告文件" in cells)
    check("含审核结论统计", "审核结论统计" in text)
    check("含出具时间", "出具时间" in cells)
    check("含审核完成时间", "审核完成时间" in cells)

    items = res.get("items") or []
    for state, title, _ in SECTIONS:
        n = sum(1 for it in items if (it.get("AI审核") or "") == state)
        if n:
            check("分级标题「%s」（%d 条）" % (title, n), title in text)
    # 每条明细：审核项名 + 依据 + 理由
    miss = [it["审核项"] for it in items
            if (it.get("AI审核") or "") in [s[0] for s in SECTIONS]
            and it["审核项"] not in all_txt]
    check("明细覆盖全部需处理条目", not miss, "缺：%s" % miss)
    check("含参考依据字段", "参考依据" in cells)
    check("含判定理由字段", "判定理由" in cells)
    # 证据行的检查要**分清两类条目**（第一次写错在这里，值得留个记录）：
    #   · 分级明细（必须整改/需核实/优化建议）：每条都带证据行（有则带页码、档位、印刷页）；
    #   · 附表（无问题/不适用）：不写证据行，但页码要出现在附表里。
    # 我原来不分青红皂白地断言"正文里必须有物理页"，结果一份报告的 15 条证据全在
    # 「无问题」项上，正文一条证据行都没有 —— 那是数据的样子，不是缺陷。
    det_states = [s[0] for s in SECTIONS]
    det_items = set(it.get("审核项") for it in items if (it.get("AI审核") or "") in det_states)
    ev = [a for a in anchors
          if (a.get("level") or "") in ("精确", "近似", "仅页码") and a.get("审核项") in det_items]
    if ev:
        check("明细证据行含页码（%d 条）" % len(ev), "物理页" in text)
        check("明细证据行写明定位档位",
              any(("证据（%s）" % lv) in text for lv in ("精确", "近似", "仅页码")),
              "没找到带档位的证据行")
        if any(a.get("printed") for a in ev):
            check("明细证据行带印刷页", "印刷页" in text)
    else:
        check("明细条目没有带页码的证据（跳过证据行检查）", True)
    # 附表里的"无问题/不适用"也要能看到证据页
    rest_ev = [a for a in anchors if a.get("page") and a.get("审核项") not in det_items]
    if rest_ev:
        check("附表给出证据页（%d 条）" % len(rest_ev), "证据页（物理/印刷）" in cells)
    check("含定位档位说明", "证据定位档位说明" in text)
    for lv in ("精确", "近似", "仅页码", "无证据"):
        check("档位说明含「%s」" % lv, lv in text)
    check("含人工复核栏", "人工复核结论" in cells)
    check("含无问题/不适用附表",
          "附表" in text if any((it.get("AI审核") or "") in ("无问题", "不适用") for it in items) else True)
    for kw in ("审核单位（盖章）", "审　核　人", "审核日期", "报告名称"):
        check("落款含「%s」" % kw, kw in text)
    check("含审核结论勾选行", "同意出具" in text and "需重新审核" in text)
    check("正文字数 > 1500", len(all_txt) > 1500, "只有 %d 字" % len(all_txt))
    check("正文非空段数 > 20", len([p for p in doc.paragraphs if p.text.strip()]) > 20)
    print("  输出：%s（%.0f KB，明细 %d 条）"
          % (os.path.basename(r["path"]), os.path.getsize(r["path"]) / 1024, r["条目"]))

print("=" * 78)
print("通过 %d / 失败 %d" % (OK, FAIL))
sys.exit(1 if FAIL else 0)
